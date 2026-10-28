#!/usr/bin/env python3
"""데모용: 같은 수정 N번을 '매번 암복호화'와 '작업구역'으로 돌려 암복호화 횟수를 비교한다.

모의 NASCA에서는 횟수만 의미가 있어 시간을 표시하지 않는다.
실제 NASCA 어댑터를 붙이면(WORKBENCH_NASCA=mcp) 같은 스크립트로 실측 시간도 잴 수 있다.

  python scripts/demo_cycles.py --edits 8
"""

from __future__ import annotations

import argparse
import io
import sys
import tempfile
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from docx import Document  # noqa: E402

from workbench.broker.manager import Manager  # noqa: E402
from workbench.config import Settings  # noqa: E402
from workbench.nasca.factory import build_nasca  # noqa: E402


class Counting:
    def __init__(self, inner):
        self.inner, self.name, self.decrypts, self.protects = inner, inner.name, 0, 0

    def decrypt(self, path):
        self.decrypts += 1
        return self.inner.decrypt(path)

    def protect(self, data, dest, *, overwrite=False):
        self.protects += 1
        return self.inner.protect(data, dest, overwrite=overwrite)

    def close(self):
        self.inner.close()


def make_source(dir_: Path, nasca) -> Path:
    d = Document()
    d.add_heading("외부 연동 설계", 1)
    for i in range(1, 9):
        d.add_paragraph(f"{i}. 항목{i}: 방식은 REST{i}로 한다.")
    buf = io.BytesIO()
    d.save(buf)
    path = dir_ / "drm" / "설계서.docx"
    nasca.inner.protect(buf.getvalue(), str(path))
    return path


def legacy(src: Path, edits: int, nasca: Counting) -> float:
    """수정할 때마다 복호화 → 수정 → 재암호화."""
    t0 = time.monotonic()
    for i in range(1, edits + 1):
        res = nasca.decrypt(str(src))
        doc = Document(io.BytesIO(res.data))
        for p in doc.paragraphs:
            if f"REST{i}" in p.text:
                p.runs[0].text = p.text.replace(f"REST{i}", f"gRPC{i}")
        buf = io.BytesIO()
        doc.save(buf)
        nasca.protect(buf.getvalue(), str(src), overwrite=True)
    return time.monotonic() - t0


def with_workbench(src: Path, edits: int, nasca: Counting, out: Path) -> float:
    settings = Settings.load()
    settings.backend = "local"
    m = Manager(settings, nasca)
    t0 = time.monotonic()
    try:
        m.open(str(src))
        for i in range(1, edits + 1):
            m.edit("설계서.docx", [{"op": "replace_text", "find": f"REST{i}", "replace": f"gRPC{i}"}])
        m.export("설계서.docx", str(out))
        m.close()
    finally:
        m.shutdown()
    return time.monotonic() - t0


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--edits", type=int, default=8)
    n = ap.parse_args().edits
    if not 1 <= n <= 8:
        ap.error("--edits는 1~8")
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        a = Counting(build_nasca())
        src = make_source(tmp, a)
        (tmp / "legacy").mkdir()
        legacy_copy = tmp / "legacy" / "설계서.docx"
        legacy_copy.write_bytes(src.read_bytes())
        t_legacy = legacy(legacy_copy, n, a)
        legacy_counts = (a.decrypts, a.protects)
        b = Counting(build_nasca())
        t_bench = with_workbench(src, n, b, tmp / "out" / "설계서_v2.docx")
        real = a.name != "mock"
        fmt = lambda t: f"{t:.2f}" if real else "-"  # noqa: E731  모의 NASCA의 시간은 의미가 없어 표시하지 않는다
        print(f"수정 {n}번 (NASCA 어댑터: {a.name})")
        print(f"{'방식':<14}{'복호화':>8}{'암호화':>8}{'시간(초)':>10}")
        print(f"{'매번 암복호화':<14}{legacy_counts[0]:>8}{legacy_counts[1]:>8}{fmt(t_legacy):>10}")
        print(f"{'작업구역':<14}{b.decrypts:>8}{b.protects:>8}{fmt(t_bench):>10}")
        if not real:
            print("모의 NASCA라서 시간은 비교하지 않는다. 실제 어댑터(WORKBENCH_NASCA=mcp)로 돌리면 시간도 나온다.")
    return 0


if __name__ == "__main__":
    sys.exit(main())

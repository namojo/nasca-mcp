"""PowerPoint(.pptx) 읽기·편집. python-pptx로 파일을 직접 다룬다."""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from ..errors import WorkbenchError
from .ooxml import blip_media_names, lossy_save_warnings


def _slide_title(slide) -> str:
    try:
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            return slide.shapes.title.text_frame.text.strip()
    except Exception:
        pass
    return ""


def _check_slide(prs, n: int):
    if not isinstance(n, int) or n < 1 or n > len(prs.slides):
        raise WorkbenchError(
            f"슬라이드 번호 {n}가 범위를 벗어났습니다. 1부터 {len(prs.slides)}까지입니다.", code="bad_index"
        )
    return prs.slides[n - 1]


def outline(path: Path) -> dict:
    prs = Presentation(str(path))
    slides = []
    for i, s in enumerate(prs.slides, start=1):
        pics = 0
        for sh in s.shapes:
            pics += len(blip_media_names(sh._element, s.part))
        slides.append({
            "slide": i,
            "title": _slide_title(s)[:80],
            "layout": s.slide_layout.name,
            "shapes": len(s.shapes),
            "pictures": pics,
            "has_notes": bool(s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip()),
        })
    return {"kind": "pptx", "slide_count": len(slides), "slides": slides,
            "layouts": [l.name for l in prs.slide_layouts],
            "warnings": lossy_save_warnings(path, "pptx")}


def _shape_lines(shape, part) -> list[str]:
    out = []
    head = f'[shape id={shape.shape_id} "{shape.name}"]'
    if shape.shape_type is not None and getattr(shape, "has_table", False) and shape.has_table:
        rows = []
        for r in shape.table.rows:
            rows.append("| " + " | ".join(c.text.replace("\n", " ⏎ ").replace("|", "\\|") for c in r.cells) + " |")
        out.append(f"{head} TABLE\n" + "\n".join(rows))
    elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        txt = shape.text_frame.text
        if txt.strip():
            out.append(f"{head} {txt}")
    media = blip_media_names(shape._element, part)
    for m in media:
        out.append(f"{head} [그림:{m}]")
    if shape.shape_type is not None and shape.shape_type == 6:  # group
        for sub in shape.shapes:
            out.extend(_shape_lines(sub, part))
    return out


def read(path: Path, *, start: int = 1, end: int | None = None, max_chars: int = 20000) -> dict:
    prs = Presentation(str(path))
    total = len(prs.slides)
    end = total if end is None else min(end, total)
    start = max(1, start)
    chunks: list[str] = []
    used = 0
    next_start = None
    for n in range(start, end + 1):
        s = prs.slides[n - 1]
        lines = [f"--- 슬라이드 {n} (레이아웃: {s.slide_layout.name}) ---"]
        for sh in s.shapes:
            lines.extend(_shape_lines(sh, s.part))
        if s.has_notes_slide:
            nt = s.notes_slide.notes_text_frame.text.strip()
            if nt:
                lines.append(f"[노트] {nt}")
        block = "\n".join(lines)
        if used + len(block) > max_chars and chunks:
            next_start = n
            break
        chunks.append(block)
        used += len(block) + 1
    return {"text": "\n".join(chunks), "total_slides": total, "next_start": next_start}


def _set_frame_text(tf, text: str) -> None:
    """첫 문단의 첫 run 서식을 유지한 채 텍스트를 바꾼다."""
    paras = tf.paragraphs
    first = paras[0]
    if first.runs:
        first.runs[0].text = text
        for r in first.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        first.add_run().text = text
    for p in paras[1:]:
        p._p.getparent().remove(p._p)


def _replace_in_frame(tf, old: str, new: str) -> int:
    count = 0
    for p in tf.paragraphs:
        for r in p.runs:
            if old in r.text:
                count += r.text.count(old)
                r.text = r.text.replace(old, new)
    return count


def _walk_text_frames(shapes):
    for sh in shapes:
        if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            yield sh.text_frame
        if getattr(sh, "has_table", False) and sh.has_table:
            for r in sh.table.rows:
                for c in r.cells:
                    yield c.text_frame
        if sh.shape_type is not None and sh.shape_type == 6:
            yield from _walk_text_frames(sh.shapes)


def _find_shape(shapes, shape_id: int):
    for sh in shapes:
        if sh.shape_id == shape_id:
            return sh
        if sh.shape_type is not None and sh.shape_type == 6:
            hit = _find_shape(sh.shapes, shape_id)
            if hit is not None:
                return hit
    return None


def _delete_slide(prs, index0: int) -> None:
    sld_id_lst = prs.slides._sldIdLst
    sld_id = list(sld_id_lst)[index0]
    prs.part.drop_rel(sld_id.rId)
    sld_id_lst.remove(sld_id)


def _apply_op(prs, op: dict, k: int, details: list) -> None:
    kind = op.get("op")
    if kind == "set_text":
        s = _check_slide(prs, op.get("slide"))
        sh = _find_shape(s.shapes, op.get("shape_id"))
        if sh is None or not getattr(sh, "has_text_frame", False) or not sh.has_text_frame:
            raise WorkbenchError(
                f"ops[{k}] 슬라이드 {op.get('slide')}에 텍스트를 가진 shape id={op.get('shape_id')}가 없습니다. "
                "workbench_read로 shape id를 확인하세요.", code="not_found")
        _set_frame_text(sh.text_frame, str(op.get("text", "")))
        details.append({"op": "set_text", "slide": op["slide"], "shape_id": op["shape_id"]})
    elif kind == "replace_text":
        find, rep = op.get("find"), op.get("replace", "")
        if not find:
            raise WorkbenchError(f"ops[{k}] replace_text에는 find가 필요합니다.", code="bad_op")
        slides = [_check_slide(prs, op["slide"])] if op.get("slide") else list(prs.slides)
        n = 0
        for s in slides:
            for tf in _walk_text_frames(s.shapes):
                n += _replace_in_frame(tf, find, rep)
        if n == 0:
            raise WorkbenchError(
                f"ops[{k}] '{find}'를 찾지 못했습니다. 서식 때문에 글자가 나뉘어 있을 수 있으니 "
                "workbench_read로 확인하세요. 아무것도 바뀌지 않았습니다.", code="not_found")
        details.append({"op": "replace_text", "replaced": n})
    elif kind == "set_notes":
        s = _check_slide(prs, op.get("slide"))
        s.notes_slide.notes_text_frame.text = str(op.get("text", ""))
        details.append({"op": "set_notes", "slide": op["slide"]})
    elif kind == "add_slide":
        layouts = list(prs.slide_layouts)
        lay = op.get("layout", 1)
        if isinstance(lay, str):
            match = [l for l in layouts if l.name == lay]
            if not match:
                raise WorkbenchError(
                    f"ops[{k}] 레이아웃 '{lay}'이 없습니다. 레이아웃: {', '.join(l.name for l in layouts)}",
                    code="bad_op")
            layout = match[0]
        else:
            if not 0 <= lay < len(layouts):
                raise WorkbenchError(f"ops[{k}] layout 번호는 0~{len(layouts) - 1}입니다.", code="bad_index")
            layout = layouts[lay]
        s = prs.slides.add_slide(layout)
        if op.get("title") and s.shapes.title is not None:
            s.shapes.title.text_frame.text = str(op["title"])
        if op.get("body"):
            for ph in s.placeholders:
                if ph.placeholder_format.idx == 1 and ph.has_text_frame:
                    ph.text_frame.text = str(op["body"])
                    break
        details.append({"op": "add_slide", "slide": len(prs.slides)})
    elif kind == "delete_slide":
        _check_slide(prs, op.get("slide"))
        _delete_slide(prs, op["slide"] - 1)
        details.append({"op": "delete_slide", "slide": op["slide"]})
    else:
        raise WorkbenchError(
            f"ops[{k}] 알 수 없는 op '{kind}'. 쓸 수 있는 op: set_text, replace_text, set_notes, add_slide, delete_slide",
            code="bad_op")


def edit(path: Path, ops: list[dict]) -> tuple[Path, dict]:
    if not ops:
        raise WorkbenchError("ops가 비어 있습니다.", code="bad_op")
    prs = Presentation(str(path))
    details: list[dict] = []
    for k, op in enumerate(ops):
        if not isinstance(op, dict):
            raise WorkbenchError(f"ops[{k}]는 객체여야 합니다.", code="bad_op")
        _apply_op(prs, op, k, details)
    fd, tmp = tempfile.mkstemp(prefix=".edit-", suffix=".pptx", dir=str(path.parent))
    os.close(fd)
    prs.save(tmp)
    return Path(tmp), {"applied": len(ops), "details": details, "slides": len(prs.slides),
                       "warnings": lossy_save_warnings(path, "pptx")}


def create(path: Path, content: dict | None) -> None:
    """content 예: {"slides": [{"title": "제목", "body": "본문"}]}"""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for item in (content or {}).get("slides", []):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        if s.shapes.title is not None:
            s.shapes.title.text_frame.text = str(item.get("title", ""))
        for ph in s.placeholders:
            if ph.placeholder_format.idx == 1 and ph.has_text_frame:
                ph.text_frame.text = str(item.get("body", ""))
    prs.save(str(path))

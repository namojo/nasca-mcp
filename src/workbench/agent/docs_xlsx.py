"""Excel(.xlsx) 읽기·편집. openpyxl로 파일을 직접 다룬다."""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.exceptions import InvalidFileException

from ..errors import WorkbenchError
from .ooxml import lossy_save_warnings

DEFAULT_ROWS = 60


def _load(path: Path, *, values: bool = False):
    try:
        return load_workbook(str(path), data_only=values)
    except InvalidFileException as e:
        raise WorkbenchError(f"엑셀 파일을 열 수 없습니다: {e}", code="bad_file") from None


def _sheet(wb, name: str | None):
    if name is None:
        return wb.worksheets[0]
    if name not in wb.sheetnames:
        raise WorkbenchError(
            f"'{name}' 시트가 없습니다. 시트: {', '.join(wb.sheetnames)}", code="not_found"
        )
    return wb[name]


def _resolve_range(ws, cell_range: str | None, *, default_rows: int | None = DEFAULT_ROWS):
    """A1:F40, A:C, 1:3 같은 표기를 (c1, r1, c2, r2)로 풀고, 사용 범위 밖이거나 거꾸로면 이유를 알려 준다."""
    max_r, max_c = max(ws.max_row, 1), max(ws.max_column, 1)
    if not cell_range:
        return 1, 1, max_c, min(max_r, default_rows) if default_rows else max_r
    try:
        c1, r1, c2, r2 = range_boundaries(cell_range)
    except ValueError:
        raise WorkbenchError(f"범위 '{cell_range}'를 해석할 수 없습니다. 예: A1:F40", code="bad_op") from None
    if c1 is None and r1 is None:
        raise WorkbenchError(f"범위 '{cell_range}'를 해석할 수 없습니다. 예: A1:F40", code="bad_op")
    c1, c2 = c1 or 1, c2 or max_c
    whole_columns = r1 is None
    r1, r2 = r1 or 1, r2 or max_r
    if r1 < 1 or c1 < 1 or r2 < r1 or c2 < c1:
        raise WorkbenchError(f"범위 '{cell_range}'가 올바르지 않습니다(왼쪽 위에서 오른쪽 아래 순서, 1부터).", code="bad_op")
    if r1 > max_r or c1 > max_c:
        raise WorkbenchError(f"범위 '{cell_range}'가 사용 범위({ws.dimensions}) 밖입니다.", code="bad_op")
    if whole_columns and default_rows:
        r2 = min(r2, r1 + default_rows - 1)
    return c1, r1, min(c2, max_c), min(r2, max_r)


def outline(path: Path) -> dict:
    wb = _load(path)
    sheets = []
    for ws in wb.worksheets:
        sheets.append({
            "name": ws.title,
            "used_range": ws.dimensions,
            "rows": ws.max_row,
            "cols": ws.max_column,
            "state": ws.sheet_state,
        })
    return {"kind": "xlsx", "sheets": sheets, "warnings": lossy_save_warnings(path, "xlsx")}


def _fmt(v) -> str:
    if v is None:
        return ""
    s = str(v)
    return s.replace("\n", "⏎").replace("|", "\\|")


def read(path: Path, *, sheet: str | None = None, cell_range: str | None = None,
         mode: str = "formulas", max_chars: int = 20000) -> dict:
    if mode not in ("formulas", "values"):
        raise WorkbenchError("mode는 'formulas' 또는 'values'입니다.", code="bad_op")
    wb = _load(path, values=(mode == "values"))
    ws = _sheet(wb, sheet)
    c1, r1, c2, r2 = _resolve_range(ws, cell_range)
    header = "    | " + " | ".join(get_column_letter(c) for c in range(c1, c2 + 1))
    lines = [f'시트 "{ws.title}" {get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2} '
             f"(사용 범위 {ws.dimensions}, mode={mode})", header]
    used = sum(len(x) for x in lines)
    next_row = None
    for r in range(r1, r2 + 1):
        row = [ws.cell(row=r, column=c).value for c in range(c1, c2 + 1)]
        line = f"{r:>3} | " + " | ".join(_fmt(v) for v in row)
        if used + len(line) > max_chars and r > r1:
            next_row = r
            break
        lines.append(line)
        used += len(line) + 1
    return {"text": "\n".join(lines), "next_row": next_row, "total_rows": ws.max_row}


def _apply_op(wb, op: dict, k: int, details: list) -> None:
    kind = op.get("op")
    if kind == "set_cells":
        ws = _sheet(wb, op.get("sheet"))
        cells = op.get("cells")
        if not isinstance(cells, dict) or not cells:
            raise WorkbenchError(f"ops[{k}] set_cells에는 cells({{\"B2\": 값}})가 필요합니다.", code="bad_op")
        for addr, val in cells.items():
            try:
                ws[addr] = val
            except Exception:
                raise WorkbenchError(f"ops[{k}] 셀 주소 '{addr}'가 올바르지 않습니다.", code="bad_op") from None
        details.append({"op": "set_cells", "sheet": ws.title, "count": len(cells)})
    elif kind == "set_range":
        ws = _sheet(wb, op.get("sheet"))
        values = op.get("values")
        top = op.get("top_left", "A1")
        if not isinstance(values, list) or not all(isinstance(r, list) for r in values):
            raise WorkbenchError(f"ops[{k}] set_range의 values는 2차원 배열이어야 합니다.", code="bad_op")
        try:
            c0, r0, _, _ = range_boundaries(top)
        except ValueError:
            raise WorkbenchError(f"ops[{k}] top_left '{top}'가 올바르지 않습니다.", code="bad_op") from None
        for dr, row in enumerate(values):
            for dc, val in enumerate(row):
                ws.cell(row=r0 + dr, column=c0 + dc, value=val)
        details.append({"op": "set_range", "sheet": ws.title, "rows": len(values)})
    elif kind == "append_rows":
        ws = _sheet(wb, op.get("sheet"))
        rows = op.get("rows")
        if not isinstance(rows, list):
            raise WorkbenchError(f"ops[{k}] append_rows에는 rows 배열이 필요합니다.", code="bad_op")
        for row in rows:
            ws.append(row)
        details.append({"op": "append_rows", "sheet": ws.title, "count": len(rows), "last_row": ws.max_row})
    elif kind == "clear_range":
        ws = _sheet(wb, op.get("sheet"))
        if not op.get("cell_range"):
            raise WorkbenchError(f"ops[{k}] clear_range에는 cell_range가 필요합니다.", code="bad_op")
        c1, r1, c2, r2 = _resolve_range(ws, op.get("cell_range"), default_rows=None)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                ws.cell(row=r, column=c).value = None
        details.append({"op": "clear_range", "sheet": ws.title})
    elif kind == "add_sheet":
        name = op.get("name")
        if not name or name in wb.sheetnames:
            raise WorkbenchError(f"ops[{k}] 시트 이름이 비었거나 이미 있습니다.", code="bad_op")
        wb.create_sheet(title=name)
        details.append({"op": "add_sheet", "name": name})
    elif kind == "rename_sheet":
        ws = _sheet(wb, op.get("sheet"))
        ws.title = op.get("name") or ws.title
        details.append({"op": "rename_sheet", "name": ws.title})
    elif kind == "delete_sheet":
        ws = _sheet(wb, op.get("sheet"))
        if len(wb.worksheets) == 1:
            raise WorkbenchError(f"ops[{k}] 마지막 시트는 지울 수 없습니다.", code="bad_op")
        wb.remove(ws)
        details.append({"op": "delete_sheet", "sheet": op.get("sheet")})
    else:
        raise WorkbenchError(
            f"ops[{k}] 알 수 없는 op '{kind}'. 쓸 수 있는 op: set_cells, set_range, append_rows, "
            "clear_range, add_sheet, rename_sheet, delete_sheet",
            code="bad_op",
        )


def edit(path: Path, ops: list[dict]) -> tuple[Path, dict]:
    if not ops:
        raise WorkbenchError("ops가 비어 있습니다.", code="bad_op")
    wb = _load(path)
    details: list[dict] = []
    for k, op in enumerate(ops):
        if not isinstance(op, dict):
            raise WorkbenchError(f"ops[{k}]는 객체여야 합니다.", code="bad_op")
        _apply_op(wb, op, k, details)
    fd, tmp = tempfile.mkstemp(prefix=".edit-", suffix=".xlsx", dir=str(path.parent))
    os.close(fd)
    wb.save(tmp)
    warnings = lossy_save_warnings(path, "xlsx")
    if any(d["op"] in ("set_cells", "set_range", "append_rows") for d in details):
        warnings.append("수식 결과는 엑셀에서 열어 다시 계산되기 전까지 저장되지 않습니다. "
                        "mode='values'로 읽으면 비어 보일 수 있습니다.")
    return Path(tmp), {"applied": len(ops), "details": details, "warnings": warnings}


def create(path: Path, content: dict | None) -> None:
    """content 예: {"sheets": {"요약": [["항목","값"],["A",1]]}}"""
    wb = Workbook()
    sheets = (content or {}).get("sheets") or {"Sheet1": []}
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        first = False
        for row in rows:
            ws.append(row)
    wb.save(str(path))

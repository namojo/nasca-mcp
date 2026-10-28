"""Word(.docx) 읽기·편집. 앱도 COM도 쓰지 않고 python-docx로 파일을 직접 다룬다."""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from ..errors import WorkbenchError
from .ooxml import blip_media_names, clip, lossy_save_warnings


def _blocks(doc):
    """본문 순서대로 (종류, 객체)를 낸다. 번호 i는 이 순서의 위치다."""
    out = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            out.append(("p", Paragraph(child, doc)))
        elif child.tag == qn("w:tbl"):
            out.append(("tbl", Table(child, doc)))
    return out


def _para_text(p: Paragraph) -> str:
    return p.text


def _style_name(p: Paragraph) -> str:
    try:
        return p.style.name if p.style is not None else ""
    except Exception:
        return ""


def _block(blocks, i: int):
    if not isinstance(i, int) or i < 0 or i >= len(blocks):
        raise WorkbenchError(
            f"블록 번호 {i}가 범위를 벗어났습니다. 0부터 {len(blocks) - 1}까지입니다. 다시 읽어 번호를 확인하세요.",
            code="bad_index",
        )
    return blocks[i]


def outline(path: Path) -> dict:
    doc = Document(str(path))
    blocks = _blocks(doc)
    heads = []
    tables = []
    chars = 0
    pics = 0
    for i, (kind, obj) in enumerate(blocks):
        if kind == "p":
            chars += len(obj.text)
            st = _style_name(obj)
            if st.lower().startswith(("heading", "제목", "title")) and obj.text.strip():
                heads.append({"i": i, "style": st, "text": obj.text.strip()[:80]})
            pics += len(blip_media_names(obj._p, doc.part))
        else:
            tables.append({"i": i, "rows": len(obj.rows), "cols": len(obj.columns)})
    return {
        "kind": "docx",
        "blocks": len(blocks),
        "chars": chars,
        "tables": tables,
        "pictures": pics,
        "headings": heads,
        "warnings": lossy_save_warnings(path, "docx"),
    }


def read(path: Path, *, start: int = 0, end: int | None = None, max_chars: int = 20000) -> dict:
    doc = Document(str(path))
    blocks = _blocks(doc)
    total = len(blocks)
    end = total if end is None else min(end, total)
    start = max(0, start)
    lines: list[str] = []
    used = 0
    next_start = None
    for i in range(start, end):
        kind, obj = blocks[i]
        if kind == "p":
            text = _para_text(obj)
            pics = blip_media_names(obj._p, doc.part)
            if not text.strip() and not pics:
                continue
            st = _style_name(obj)
            tag = f" ({st})" if st and st not in ("Normal", "표준", "본문") else ""
            pic = "".join(f" [그림:{n}]" for n in pics)
            line = f"[{i}]{tag} {text}{pic}"
        else:
            rows = []
            for r in obj.rows:
                cells = []
                seen = set()
                for c in r.cells:
                    if id(c._tc) in seen:
                        continue
                    seen.add(id(c._tc))
                    cells.append(c.text.replace("\n", " ⏎ ").replace("|", "\\|"))
                rows.append("| " + " | ".join(cells) + " |")
            line = f"[{i}] TABLE {len(obj.rows)}x{len(obj.columns)}\n" + "\n".join(rows)
        if used + len(line) > max_chars and lines:
            next_start = i
            break
        lines.append(line)
        used += len(line) + 1
    return {"text": "\n".join(lines), "total_blocks": total, "next_start": next_start}


# --- 편집 ---------------------------------------------------------------

def _replace_in_paragraph(p: Paragraph, old: str, new: str) -> int:
    """문단 안의 old를 new로 바꾼다. 여러 run에 걸쳐 있어도 첫 run의 서식으로 합친다."""
    count = 0
    search_from = 0
    while True:
        runs = p.runs
        full = "".join(r.text for r in runs)
        pos = full.find(old, search_from)
        if pos < 0:
            return count
        end = pos + len(old)
        acc = 0
        first_idx = last_idx = None
        first_off = last_off = 0
        for idx, r in enumerate(runs):
            ln = len(r.text)
            if first_idx is None and pos < acc + ln:
                first_idx, first_off = idx, pos - acc
            if first_idx is not None and end <= acc + ln:
                last_idx, last_off = idx, end - acc
                break
            acc += ln
        if first_idx is None or last_idx is None:
            return count
        if first_idx == last_idx:
            r = runs[first_idx]
            r.text = r.text[:first_off] + new + r.text[last_off:]
        else:
            runs[first_idx].text = runs[first_idx].text[:first_off] + new
            for k in range(first_idx + 1, last_idx):
                runs[k].text = ""
            runs[last_idx].text = runs[last_idx].text[last_off:]
        count += 1
        search_from = pos + len(new)


def _all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                for p in cell.paragraphs:
                    yield p


def _set_paragraph_text(p: Paragraph, text: str) -> None:
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def _resolve_style(doc, style: str):
    try:
        return doc.styles[style]
    except KeyError:
        names = sorted(s.name for s in doc.styles if getattr(s, "type", None) == 1)[:25]
        raise WorkbenchError(
            f"문서에 '{style}' 스타일이 없습니다. 쓸 수 있는 문단 스타일 예: {', '.join(names)}",
            code="bad_style",
        ) from None


def _apply_op(doc, op: dict, k: int, details: list) -> None:
    kind = op.get("op")
    if kind == "replace_text":
        find, rep = op.get("find"), op.get("replace", "")
        if not find:
            raise WorkbenchError(f"ops[{k}] replace_text에는 find가 필요합니다.", code="bad_op")
        n = sum(_replace_in_paragraph(p, find, rep) for p in _all_paragraphs(doc))
        if n == 0:
            raise WorkbenchError(
                f"ops[{k}] '{find}'를 찾지 못했습니다. 문서 서식 때문에 글자가 나뉘어 있을 수 있으니 "
                "workbench_read로 실제 문구를 확인하세요. 아무것도 바뀌지 않았습니다.",
                code="not_found",
            )
        details.append({"op": "replace_text", "replaced": n})
        return

    blocks = _blocks(doc)
    if kind == "set_text":
        t, obj = _block(blocks, op.get("i"))
        if t != "p":
            raise WorkbenchError(f"ops[{k}] 블록 {op['i']}는 표입니다. set_cell을 쓰세요.", code="bad_op")
        _set_paragraph_text(obj, str(op.get("text", "")))
        details.append({"op": "set_text", "i": op["i"]})
    elif kind in ("insert_after", "insert_before"):
        t, anchor = _block(blocks, op.get("i"))
        style = op.get("style")
        new_p = doc.add_paragraph(str(op.get("text", "")))
        if style:
            new_p.style = _resolve_style(doc, style)
        el = new_p._p
        el.getparent().remove(el)
        target = anchor._p if t == "p" else anchor._tbl
        (target.addnext if kind == "insert_after" else target.addprevious)(el)
        details.append({"op": kind, "i": op["i"]})
    elif kind == "delete":
        t, obj = _block(blocks, op.get("i"))
        el = obj._p if t == "p" else obj._tbl
        el.getparent().remove(el)
        details.append({"op": "delete", "i": op["i"]})
    elif kind == "set_style":
        t, obj = _block(blocks, op.get("i"))
        if t != "p":
            raise WorkbenchError(f"ops[{k}] 블록 {op['i']}는 표입니다.", code="bad_op")
        obj.style = _resolve_style(doc, op.get("style", ""))
        details.append({"op": "set_style", "i": op["i"]})
    elif kind == "set_cell":
        t, tbl = _block(blocks, op.get("i"))
        if t != "tbl":
            raise WorkbenchError(f"ops[{k}] 블록 {op['i']}는 표가 아닙니다.", code="bad_op")
        try:
            cell = tbl.cell(op["row"], op["col"])
        except (IndexError, KeyError):
            raise WorkbenchError(
                f"ops[{k}] 표는 {len(tbl.rows)}행 {len(tbl.columns)}열입니다. row/col은 0부터입니다.",
                code="bad_index",
            ) from None
        paras = cell.paragraphs
        _set_paragraph_text(paras[0], str(op.get("text", "")))
        for extra in paras[1:]:
            extra._p.getparent().remove(extra._p)
        details.append({"op": "set_cell", "i": op["i"], "row": op["row"], "col": op["col"]})
    elif kind == "add_row":
        t, tbl = _block(blocks, op.get("i"))
        if t != "tbl":
            raise WorkbenchError(f"ops[{k}] 블록 {op['i']}는 표가 아닙니다.", code="bad_op")
        row = tbl.add_row()
        for c, val in zip(row.cells, op.get("cells", [])):
            _set_paragraph_text(c.paragraphs[0], str(val))
        details.append({"op": "add_row", "i": op["i"], "rows": len(tbl.rows)})
    else:
        raise WorkbenchError(
            f"ops[{k}] 알 수 없는 op '{kind}'. 쓸 수 있는 op: replace_text, set_text, insert_after, "
            "insert_before, delete, set_style, set_cell, add_row",
            code="bad_op",
        )


def edit(path: Path, ops: list[dict]) -> tuple[Path, dict]:
    """ops를 순서대로 적용해 임시 파일에 저장하고 (임시 경로, 결과)를 돌려준다. 실패하면 아무것도 바뀌지 않는다."""
    if not ops:
        raise WorkbenchError("ops가 비어 있습니다.", code="bad_op")
    doc = Document(str(path))
    details: list[dict] = []
    for k, op in enumerate(ops):
        if not isinstance(op, dict):
            raise WorkbenchError(f"ops[{k}]는 객체여야 합니다.", code="bad_op")
        _apply_op(doc, op, k, details)
    fd, tmp = tempfile.mkstemp(prefix=".edit-", suffix=".docx", dir=str(path.parent))
    os.close(fd)
    doc.save(tmp)
    return Path(tmp), {"applied": len(ops), "details": details, "blocks": len(_blocks(doc)),
                       "warnings": lossy_save_warnings(path, "docx")}


def create(path: Path, content: list | None) -> None:
    doc = Document()
    for item in content or []:
        if isinstance(item, str):
            doc.add_paragraph(item)
        elif isinstance(item, dict) and "heading" in item:
            doc.add_heading(str(item["heading"]), level=int(item.get("level", 1)))
        elif isinstance(item, dict) and "table" in item:
            rows = item["table"]
            if rows:
                t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                t.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c, val in enumerate(row):
                        t.cell(r, c).text = str(val)
        elif isinstance(item, dict) and "p" in item:
            para = doc.add_paragraph(str(item["p"]))
            if item.get("style"):
                para.style = _resolve_style(doc, item["style"])
        else:
            raise WorkbenchError(
                "content 항목은 문자열, {heading,level}, {p,style}, {table:[[..]]} 중 하나여야 합니다.",
                code="bad_op",
            )
    doc.save(str(path))

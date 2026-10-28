import io
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from workbench.broker.manager import Manager  # noqa: E402
from workbench.config import Settings  # noqa: E402
from workbench.nasca.mock import MockNasca  # noqa: E402


def make_docx(path: Path, paragraphs=("외부 연동 API는 REST로 한다.", "인증은 OAuth2를 쓴다.")):
    from docx import Document

    d = Document()
    d.add_heading("1. 개요", level=1)
    for p in paragraphs:
        d.add_paragraph(p)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "API", "방식"
    t.cell(1, 0).text, t.cell(1, 1).text = "주문", "REST"
    d.save(str(path))


def make_png(path: Path, size=(1200, 800), color="blue"):
    from PIL import Image

    Image.new("RGB", size, color).save(str(path), "PNG")


@pytest.fixture
def settings(tmp_path):
    s = Settings()
    s.home = tmp_path / "home"
    s.backend = "local"
    s.nasca = "mock"
    s.idle_ttl_s = 60
    s.max_ttl_s = 600
    s.lease_s = 30
    return s


@pytest.fixture
def nasca():
    return MockNasca()


@pytest.fixture
def manager(settings, nasca):
    m = Manager(settings, nasca)
    yield m
    m.shutdown()


@pytest.fixture
def protected_docx(tmp_path, nasca):
    """모의 NASCA로 보호한 docx. 일반 읽기로는 문서로 열리지 않는다."""
    plain = tmp_path / "plain.docx"
    make_docx(plain)
    out = tmp_path / "drm" / "설계서.docx"
    nasca.protect(plain.read_bytes(), str(out))
    return out


@pytest.fixture
def protected_png(tmp_path, nasca):
    plain = tmp_path / "plain.png"
    make_png(plain)
    out = tmp_path / "drm" / "구성도.png"
    nasca.protect(plain.read_bytes(), str(out))
    return out

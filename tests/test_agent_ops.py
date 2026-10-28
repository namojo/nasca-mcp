import base64
import io
import zipfile

import pytest
from PIL import Image

from workbench.agent.rpc import Dispatcher
from workbench.agent.store import Store
from workbench.errors import WorkbenchError


@pytest.fixture
def d(tmp_path):
    return Dispatcher(Store(tmp_path / "root"))


def put(d, name, data: bytes):
    return d.call("put", {"name": name, "data_b64": base64.b64encode(data).decode()})


def png_bytes(size=(300, 200), color="red", fmt="PNG"):
    b = io.BytesIO()
    Image.new("RGB", size, color).save(b, fmt)
    return b.getvalue()


# --- docx ---------------------------------------------------------------
def test_docx_replace_across_runs_keeps_first_run_format(d, tmp_path):
    from docx import Document

    doc = Document()
    p = doc.add_paragraph()
    p.add_run("외부 연동은 RE").bold = True
    p.add_run("ST 방식").bold = False
    f = tmp_path / "split.docx"
    doc.save(str(f))
    put(d, "split.docx", f.read_bytes())
    r = d.call("edit", {"name": "split.docx", "ops": [{"op": "replace_text", "find": "REST", "replace": "gRPC"}]})
    assert r["details"][0]["replaced"] == 1
    assert "외부 연동은 gRPC 방식" in d.call("read", {"name": "split.docx"})["text"]


def test_docx_failed_op_changes_nothing(d):
    d.call("create", {"name": "a.docx", "content": ["하나", "둘"]})
    before = d.call("get", {"name": "a.docx"})["sha256"]
    with pytest.raises(WorkbenchError, match="찾지 못했습니다"):
        d.call("edit", {"name": "a.docx", "ops": [
            {"op": "set_text", "i": 0, "text": "바뀜"},
            {"op": "replace_text", "find": "없음", "replace": "x"},
        ]})
    assert d.call("get", {"name": "a.docx"})["sha256"] == before, "ops는 전부 적용되거나 전부 취소돼야 한다"


def test_docx_bad_index_and_style_messages_guide_the_model(d):
    d.call("create", {"name": "a.docx", "content": ["하나"]})
    with pytest.raises(WorkbenchError, match="범위를 벗어났습니다"):
        d.call("edit", {"name": "a.docx", "ops": [{"op": "set_text", "i": 9, "text": "x"}]})
    with pytest.raises(WorkbenchError, match="스타일이 없습니다"):
        d.call("edit", {"name": "a.docx", "ops": [{"op": "set_style", "i": 0, "style": "없는스타일"}]})
    with pytest.raises(WorkbenchError, match="쓸 수 있는 op"):
        d.call("edit", {"name": "a.docx", "ops": [{"op": "explode"}]})


def test_docx_table_ops_and_delete_and_undo(d):
    d.call("create", {"name": "t.docx", "content": ["앞", {"table": [["A", "B"], ["1", "2"]]}, "뒤"]})
    d.call("edit", {"name": "t.docx", "ops": [{"op": "set_cell", "i": 1, "row": 1, "col": 1, "text": "9"},
                                              {"op": "add_row", "i": 1, "cells": ["3", "4"]}]})
    text = d.call("read", {"name": "t.docx"})["text"]
    assert "| 1 | 9 |" in text and "| 3 | 4 |" in text and "TABLE 3x2" in text
    d.call("edit", {"name": "t.docx", "ops": [{"op": "delete", "i": 0}]})
    assert "[0] TABLE" in d.call("read", {"name": "t.docx"})["text"]
    d.call("revert", {"name": "t.docx"})
    assert "[0] 앞" in d.call("read", {"name": "t.docx"})["text"]


def test_docx_read_is_paged(d):
    d.call("create", {"name": "long.docx", "content": [f"문단 {i} " + "가" * 200 for i in range(60)]})
    first = d.call("read", {"name": "long.docx", "max_chars": 2000})
    assert first["next_start"] is not None and first["total_blocks"] >= 60
    rest = d.call("read", {"name": "long.docx", "start": first["next_start"], "max_chars": 100000})
    assert rest["next_start"] is None and "문단 59" in rest["text"]


def test_docx_picture_is_marked_and_extractable(d, tmp_path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("그림 앞")
    img = tmp_path / "pic.png"
    img.write_bytes(png_bytes())
    doc.add_picture(str(img))
    f = tmp_path / "withpic.docx"
    doc.save(str(f))
    put(d, "withpic.docx", f.read_bytes())
    text = d.call("read", {"name": "withpic.docx"})["text"]
    assert "[그림:image1.png]" in text
    assert d.call("outline", {"name": "withpic.docx"})["pictures"] == 1
    saved = d.call("extract_media", {"name": "withpic.docx"})["saved"]
    assert saved == ["withpic.docx.media/image1.png"]
    im = d.call("image", {"name": saved[0]})
    assert im["original_width"] == 300


# --- xlsx ---------------------------------------------------------------
def test_xlsx_formulas_and_values_modes_and_range(d):
    d.call("create", {"name": "s.xlsx", "content": {"sheets": {"요약": [["항목", "값"], ["A", 1], ["B", 2]]}}})
    d.call("edit", {"name": "s.xlsx", "ops": [{"op": "set_cells", "cells": {"A4": "합계", "B4": "=B2+B3"}}]})
    assert "=B2+B3" in d.call("read", {"name": "s.xlsx"})["text"]
    sub = d.call("read", {"name": "s.xlsx", "cell_range": "A2:B3"})["text"]
    assert "항목" not in sub and "| A | 1" in sub
    with pytest.raises(WorkbenchError, match="시트가 없습니다"):
        d.call("read", {"name": "s.xlsx", "sheet": "없음"})
    with pytest.raises(WorkbenchError, match="해석할 수 없습니다"):
        d.call("read", {"name": "s.xlsx", "cell_range": "A1:F"})
    with pytest.raises(WorkbenchError, match="사용 범위"):
        d.call("read", {"name": "s.xlsx", "cell_range": "zzz"})
    with pytest.raises(WorkbenchError, match="올바르지 않습니다"):
        d.call("read", {"name": "s.xlsx", "cell_range": "B2:A1"})
    assert "| 값" in d.call("read", {"name": "s.xlsx", "cell_range": "B:B"})["text"]
    d.call("edit", {"name": "s.xlsx", "ops": [{"op": "clear_range", "cell_range": "A2:B3"}]})
    assert "| A | 1" not in d.call("read", {"name": "s.xlsx"})["text"]


def test_xlsx_warns_when_charts_would_be_lost(d, tmp_path):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference

    wb = Workbook()
    ws = wb.active
    for r in ([1], [2], [3]):
        ws.append(r)
    ch = BarChart()
    ch.add_data(Reference(ws, min_col=1, min_row=1, max_row=3))
    ws.add_chart(ch, "D2")
    f = tmp_path / "chart.xlsx"
    wb.save(str(f))
    put(d, "chart.xlsx", f.read_bytes())
    assert any("차트" in w or "xl/charts" in w for w in d.call("outline", {"name": "chart.xlsx"})["warnings"])
    r = d.call("edit", {"name": "chart.xlsx", "ops": [{"op": "set_cells", "cells": {"A1": 5}}]})
    assert any("xl/charts" in w for w in r["warnings"])


# --- pptx ---------------------------------------------------------------
def test_pptx_edit_by_shape_id_notes_and_delete(d):
    d.call("create", {"name": "p.pptx", "content": {"slides": [{"title": "하나", "body": "본문1"},
                                                              {"title": "둘", "body": "본문2"}]}})
    text = d.call("read", {"name": "p.pptx"})["text"]
    assert "슬라이드 1" in text and "shape id=3" in text
    d.call("edit", {"name": "p.pptx", "ops": [{"op": "set_text", "slide": 1, "shape_id": 3, "text": "바뀐 본문"},
                                             {"op": "set_notes", "slide": 2, "text": "발표 노트"},
                                             {"op": "delete_slide", "slide": 1}]})
    out = d.call("read", {"name": "p.pptx"})["text"]
    assert "바뀐 본문" not in out and "[노트] 발표 노트" in out
    assert d.call("outline", {"name": "p.pptx"})["slide_count"] == 1
    with pytest.raises(WorkbenchError, match="shape id"):
        d.call("edit", {"name": "p.pptx", "ops": [{"op": "set_text", "slide": 1, "shape_id": 999, "text": "x"}]})


# --- 이미지 -------------------------------------------------------------
def test_image_render_limits_and_crop(d):
    put(d, "big.png", png_bytes((4000, 2000)))
    r = d.call("image", {"name": "big.png", "max_side": 1000})
    assert (r["width"], r["height"]) == (1000, 500)
    r = d.call("image", {"name": "big.png", "crop_px": [0, 0, 400, 200]})
    assert (r["width"], r["height"]) == (400, 200)
    with pytest.raises(WorkbenchError, match="하나만"):
        d.call("image", {"name": "big.png", "crop_px": [0, 0, 1, 1], "crop_frac": [0, 0, .5, .5]})
    with pytest.raises(WorkbenchError):
        d.call("image", {"name": "big.png", "crop_frac": [0.6, 0, 0.4, 1]})


def test_image_edit_keeps_format_and_can_undo(d):
    put(d, "a.jpg", png_bytes((800, 400), fmt="JPEG"))
    d.call("edit", {"name": "a.jpg", "ops": [{"op": "rotate", "degrees": 90}, {"op": "resize", "max_side": 300}]})
    info = d.call("outline", {"name": "a.jpg"})
    assert info["format"] == "JPEG" and max(info["width"], info["height"]) == 300
    d.call("revert", {"name": "a.jpg"})
    assert d.call("outline", {"name": "a.jpg"})["width"] == 800


def test_corrupt_image_is_reported_not_crashed(d):
    r = put(d, "bad.png", b"not an image")
    assert "outline_error" in r
    with pytest.raises(WorkbenchError, match="열 수 없습니다"):
        d.call("image", {"name": "bad.png"})


# --- 텍스트 -------------------------------------------------------------
def test_text_preserves_cp949(d):
    put(d, "m.txt", "가나다\n라마바\n".encode("cp949"))
    d.call("edit", {"name": "m.txt", "ops": [{"op": "replace_text", "find": "라마바", "replace": "사아자"}]})
    raw = d.call("get", {"name": "m.txt"})["data_b64"]
    assert base64.b64decode(raw).decode("cp949") == "가나다\n사아자\n"


def test_text_line_ops(d):
    d.call("create", {"name": "n.txt", "content": "1\n2\n3\n"})
    d.call("edit", {"name": "n.txt", "ops": [{"op": "delete_lines", "start": 2, "end": 2},
                                            {"op": "insert_lines", "after_line": 0, "text": "0"}]})
    lines = [l.split(None, 1)[1] for l in d.call("read", {"name": "n.txt"})["text"].splitlines()]
    assert lines == ["0", "1", "3"]


# --- 저장소 안전 --------------------------------------------------------
@pytest.mark.parametrize("bad", ["../x.txt", "/etc/passwd", "C:\\Windows\\a.txt", "a/../../b.txt", "con.txt",
                                 ".rev/x", "a:b.txt", "x" * 300, "", "  "])
def test_unsafe_names_are_rejected(d, bad):
    with pytest.raises(WorkbenchError):
        put(d, bad, b"x")


def test_size_limits(tmp_path):
    d = Dispatcher(Store(tmp_path / "r", max_file_bytes=100))
    with pytest.raises(WorkbenchError, match="너무 큽니다"):
        put(d, "a.txt", b"x" * 101)


def test_delete_shreds_file_and_revisions(d, tmp_path):
    d.call("create", {"name": "a.txt", "content": "비밀"})
    d.call("edit", {"name": "a.txt", "ops": [{"op": "append", "text": "더"}]})
    root = d.store.root
    d.call("delete", {"name": "a.txt"})
    assert [p for p in root.rglob("*") if p.is_file()] == []
    assert d.call("list", {})["files"] == []


def test_wipe_removes_everything(d):
    d.call("create", {"name": "a.txt", "content": "x"})
    d.call("edit", {"name": "a.txt", "ops": [{"op": "append", "text": "y"}]})
    d.call("wipe", {})
    assert not [p for p in d.store.root.rglob("*") if p.is_file()]


def test_unsupported_format_listed_but_not_readable(d):
    put(d, "a.hwp", b"\xd0\xcf\x11\xe0")
    with pytest.raises(WorkbenchError, match="지원하지 않는"):
        d.call("read", {"name": "a.hwp"})
    assert d.call("get", {"name": "a.hwp"})["size"] == 4
